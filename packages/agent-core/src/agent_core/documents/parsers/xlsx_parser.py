"""XLSX parser (Phase 4) — openpyxl behind the parser interface.

openpyxl is an *optional* dependency (``agent-core[docs]``) imported
lazily. One section per sheet, preserving sheet names and the cell grid as
a deterministic textual representation (rows joined by newlines, cells by
`` | ``). Workbooks are opened read-only with ``data_only=True`` (cached
values only — formulas are never evaluated, so embedded content cannot run).
"""

from __future__ import annotations

import contextlib
import io

from ..errors import (
    DOCUMENT_CORRUPT,
    EXTRACTION_FAILED,
    PARSER_UNAVAILABLE,
    DocumentError,
)
from ..limits import DocumentLimits
from .base import SECTION_SHEET, DocumentParser, ExtractedContent, RawSection

MAX_ROWS_PER_SHEET = 10_000
MAX_COLS_PER_SHEET = 200


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


class XlsxParser(DocumentParser):
    name = "xlsx"
    document_type = "xlsx"
    supported_extensions = ("xlsx",)
    supported_media_types = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",)

    def _extract(self, data: bytes, limits: DocumentLimits) -> ExtractedContent:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - depends on install extras
            raise DocumentError(
                PARSER_UNAVAILABLE,
                "XLSX parsing requires the 'openpyxl' package (install agent-core[docs])",
            ) from exc

        try:
            workbook = load_workbook(
                io.BytesIO(data), read_only=True, data_only=True, keep_links=False
            )
        except Exception as exc:
            raise DocumentError(
                DOCUMENT_CORRUPT, f"XLSX could not be parsed: {type(exc).__name__}"
            ) from exc

        try:
            sections: list[RawSection] = []
            warnings: list[str] = []
            truncated = False
            max_sheets = limits.max_sheets
            sheet_names = list(workbook.sheetnames)
            total_sheets = len(sheet_names)
            if len(sheet_names) > max_sheets:
                truncated = True
                warnings.append(
                    f"only the first {max_sheets} sheets were extracted "
                    f"({len(sheet_names) - max_sheets} sheet(s) skipped)"
                )
            for name in sheet_names[:max_sheets]:
                sheet = workbook[name]
                rows: list[list[str]] = []
                col_count = 0
                for row in sheet.iter_rows(min_col=1, max_col=MAX_COLS_PER_SHEET):
                    values = [_cell_to_text(cell.value) for cell in row]
                    while values and values[-1] == "":
                        values.pop()  # trailing empty cells are grid padding, not data
                    if any(values):
                        rows.append(values)
                        col_count = max(col_count, len(values))
                    if len(rows) >= MAX_ROWS_PER_SHEET:
                        warnings.append(f"sheet {name!r}: rows truncated at {MAX_ROWS_PER_SHEET}")
                        truncated = True
                        break
                lines = [" | ".join(r) for r in rows]
                sections.append(
                    RawSection(
                        section_type=SECTION_SHEET,
                        text="\n".join(lines),
                        location={"sheet": name},
                        metadata={"rows": len(rows), "columns": col_count},
                    )
                )
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError(
                EXTRACTION_FAILED, f"XLSX extraction failed: {type(exc).__name__}"
            ) from exc
        finally:
            with contextlib.suppress(Exception):
                workbook.close()

        return ExtractedContent(
            document_type=self.document_type,
            media_type=self.supported_media_types[0],
            title=None,
            sections=sections,
            warnings=warnings,
            stats={
                "sheets": total_sheets,
                "sheets_extracted": min(len(sheet_names), limits.max_sheets),
            },
            truncated=truncated,
        )

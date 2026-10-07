"""PPTX parser (Phase 4) — python-pptx behind the parser interface.

python-pptx is an *optional* dependency (``agent-core[docs]``) imported
lazily. One section per slide, preserving slide numbers and the in-slide
text ordering of shapes. python-pptx reads structure only and never runs
embedded content.
"""

from __future__ import annotations

import io

from ..errors import (
    DOCUMENT_CORRUPT,
    EXTRACTION_FAILED,
    PARSER_UNAVAILABLE,
    DocumentError,
)
from ..limits import DocumentLimits
from .base import SECTION_SLIDE, DocumentParser, ExtractedContent, RawSection


class PptxParser(DocumentParser):
    name = "pptx"
    document_type = "pptx"
    supported_extensions = ("pptx",)
    supported_media_types = (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )

    def _extract(self, data: bytes, limits: DocumentLimits) -> ExtractedContent:
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover - depends on install extras
            raise DocumentError(
                PARSER_UNAVAILABLE,
                "PPTX parsing requires the 'python-pptx' package (install agent-core[docs])",
            ) from exc

        try:
            prs = Presentation(io.BytesIO(data))
            slides = list(prs.slides)
        except Exception as exc:
            raise DocumentError(
                DOCUMENT_CORRUPT, f"PPTX could not be parsed: {type(exc).__name__}"
            ) from exc

        try:
            sections: list[RawSection] = []
            warnings: list[str] = []
            truncated = False
            max_slides = limits.max_slides
            if len(slides) > max_slides:
                truncated = True
                warnings.append(
                    f"only the first {max_slides} slides were extracted "
                    f"({len(slides) - max_slides} slide(s) skipped)"
                )
            title: str | None = None
            for i in range(min(len(slides), max_slides)):
                slide = slides[i]
                texts: list[str] = []
                slide_title: str | None = None
                try:
                    for shape in slide.shapes:
                        if not shape.has_text_frame:
                            continue
                        text = shape.text_frame.text.strip()
                        if not text:
                            continue
                        if slide_title is None and getattr(shape, "is_placeholder", False):
                            ph = shape.placeholder_format
                            if ph is not None and ph.idx == 0:
                                slide_title = text
                        texts.append(text)
                except Exception as exc:
                    warnings.append(f"slide {i + 1}: extraction failed ({type(exc).__name__})")
                body = "\n".join(texts)
                if title is None and slide_title:
                    title = slide_title
                sections.append(
                    RawSection(
                        section_type=SECTION_SLIDE,
                        text=body,
                        heading=slide_title,
                        location={"slide": i + 1},
                    )
                )
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError(
                EXTRACTION_FAILED, f"PPTX text extraction failed: {type(exc).__name__}"
            ) from exc

        return ExtractedContent(
            document_type=self.document_type,
            media_type=self.supported_media_types[0],
            title=title,
            sections=sections,
            warnings=warnings,
            stats={"slides": len(slides), "slides_extracted": min(len(slides), limits.max_slides)},
            truncated=truncated,
        )

"""Deterministic document fixture builders for Phase 4 tests.

Every fixture is built in a temp directory from constant content — no real
user files, no network, no external binaries. The PDF is generated
structurally (a minimal valid multi-page PDF) so the same bytes parse the
same way on every run.
"""

from __future__ import annotations

from pathlib import Path


def make_pdf(path: Path, pages: list[str], title: str | None = None) -> Path:
    """Write a minimal valid multi-page PDF (pypdf-readable).

    Layout: Helvetica text stream per page. Deterministic object numbering
    and xref offsets; pypdf extracts the page text exactly as given.
    """
    objects: list[bytes] = []

    def add(obj: bytes) -> int:
        objects.append(obj)
        return len(objects)

    n = len(pages)
    cat = add(b"")
    pages_id = add(b"")
    content_ids = [add(b"") for _ in pages]
    page_ids = [add(b"") for _ in pages]
    font = add(b"")

    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects[pages_id - 1] = (
        f"{pages_id} 0 obj\n<< /Type /Pages /Kids [{kids}] /Count {n} >>\nendobj\n"
    ).encode()
    objects[cat - 1] = (
        f"{cat} 0 obj\n<< /Type /Catalog /Pages {pages_id} 0 R >>\nendobj\n"
    ).encode()
    for i, text in enumerate(pages):
        safe = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        stream = f"BT /F1 12 Tf 50 750 Td ({safe}) Tj ET".encode()
        objects[content_ids[i] - 1] = (
            f"{content_ids[i]} 0 obj\n<< /Length {len(stream)} >>\nstream\n".encode()
            + stream
            + b"\nendstream\nendobj\n"
        )
        objects[page_ids[i] - 1] = (
            f"{page_ids[i]} 0 obj\n<< /Type /Page /Parent {pages_id} 0 R "
            f"/MediaBox [0 0 612 792] /Resources << /Font << /F1 {font} 0 R >> >> "
            f"/Contents {content_ids[i]} 0 R >>\nendobj\n"
        ).encode()
    objects[font - 1] = (
        f"{font} 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
    ).encode()

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for obj in objects:
        offsets.append(len(out))
        out += obj
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root {cat} 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    path.write_bytes(bytes(out))
    return path


def make_docx(path: Path) -> Path:
    """A DOCX with a title, headings, paragraphs, and a 2x2 table."""
    from docx import Document

    doc = Document()
    doc.add_heading("Quarterly Report", level=0)  # style: Title
    doc.add_heading("Revenue", level=1)
    doc.add_paragraph("The revenue grew by twelve percent this quarter.")
    doc.add_heading("Expenses", level=2)
    doc.add_paragraph("Operating expenses stayed flat.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "12%"
    doc.save(str(path))
    return path


def make_pptx(path: Path) -> Path:
    """A PPTX with two slides: title + body text each."""
    from pptx import Presentation

    prs = Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[1])
    slide1.shapes.title.text = "Slide One"
    slide1.placeholders[1].text = "Quarterly summary goes here."
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Slide Two"
    slide2.placeholders[1].text = "Budget numbers on the next slide."
    prs.save(str(path))
    return path


def make_xlsx(path: Path) -> Path:
    """An XLSX with two sheets and a small cell grid each."""
    from openpyxl import Workbook

    wb = Workbook()
    first = wb.active
    first.title = "Data"
    first["A1"], first["B1"] = "Product", "Units"
    first["A2"], first["B2"] = "Widget", 30
    second = wb.create_sheet("Notes")
    second["A1"] = "Shipped in Q3"
    wb.save(str(path))
    return path

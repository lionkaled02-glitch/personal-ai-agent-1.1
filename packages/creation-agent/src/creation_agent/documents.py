"""Safe document generation adapters with bounded inputs."""

from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path


class DocumentGenerationError(RuntimeError):
    pass


def _safe_name(name: str, suffix: str) -> str:
    clean = Path(name).name.replace("\x00", "")
    if not clean:
        clean = "document" + suffix
    if not clean.lower().endswith(suffix):
        clean += suffix
    return clean[:180]


def generate_docx(title: str, paragraphs: Iterable[str], output: Path) -> Path:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentGenerationError("python-docx is required") from exc
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title[:500], level=1)
    for paragraph in list(paragraphs)[:500]:
        doc.add_paragraph(str(paragraph)[:20_000])
    doc.save(output)
    return output


def generate_xlsx(title: str, rows: Iterable[Iterable[object]], output: Path) -> Path:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise DocumentGenerationError("openpyxl is required") from exc
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Sheet1"
    for row in list(rows)[:10_000]:
        ws.append([str(v)[:10_000] if v is not None else None for v in list(row)[:100]])
    wb.save(output)
    return output


def generate_pptx(title: str, slides: Iterable[tuple[str, str]], output: Path) -> Path:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise DocumentGenerationError("python-pptx is required") from exc
    output.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    for heading, body in list(slides)[:100]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str(heading)[:500]
        slide.placeholders[1].text = str(body)[:20_000]
    prs.save(output)
    return output


def generate_pdf(title: str, paragraphs: Iterable[str], output: Path) -> Path:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise DocumentGenerationError("reportlab is required") from exc
    output.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    story = [Paragraph(title[:500], styles["Title"]), Spacer(1, 12)]
    for paragraph in list(paragraphs)[:500]:
        story.extend([Paragraph(str(paragraph)[:20_000], styles["BodyText"]), Spacer(1, 8)])
    SimpleDocTemplate(str(output), pagesize=A4).build(story)
    return output
